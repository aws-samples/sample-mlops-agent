"""MLflow skill Lambda — Gateway MCP target.

Tools:
  - list_scorers                — dynamic catalog of MLflow built-in judge
    scorers available to submit_eval_job. The agent calls this to translate
    user-provided scorer terms (e.g. "faithfulness", "answer relevance")
    into the canonical names the SageMaker eval container expects.
  - query_metrics
  - retrieve_traces
  - analyze_trace
  - generate_compliance_report  — renders an eval run into markdown + .docx,
    uploads both to S3, stamps the DynamoDB thread row with the artifact URI.
    The agent authors the narrative; this tool persists it deterministically.
"""
import io
import json
import os
import tempfile
import time
from typing import Any

import boto3
import mlflow

PROJECT_NAME        = os.environ.get("PROJECT_NAME", "sample-mlops-agent")
JOBS_TABLE          = os.environ.get("JOBS_TABLE", f"{PROJECT_NAME}-metadata")
SESSION_BUCKET      = os.environ.get("SESSION_BUCKET", "")
AWS_REGION          = os.environ.get("AWS_REGION", "us-east-1")
MLFLOW_TRACKING_URI = os.environ.get("MLFLOW_TRACKING_URI", "")

# Failing-rows appendix defaults — worst N rows by aggregate scorer mean.
_APPENDIX_ROWS = 10


def _mlflow_client() -> mlflow.tracking.MlflowClient:
    """Return configured MlflowClient pointed at SageMaker MLflow App."""
    if MLFLOW_TRACKING_URI:
        mlflow.set_tracking_uri(MLFLOW_TRACKING_URI)
    return mlflow.tracking.MlflowClient()


def _list_scorers(args: dict) -> dict:
    """Return the live catalog of MLflow built-in judge scorers.

    Introspects `mlflow.genai.scorers` at runtime (no hard-coded list) so
    newly released scorers appear automatically on the next deploy of the
    mlflow skill image. The SageMaker eval container's SCORER_CATALOG is a
    subset of this list; the agent should consult this tool BEFORE calling
    submit_eval_job when the user provides scorer terms in plain English.

    The agent is responsible for the semantic match between user vocabulary
    ("faithfulness", "answer relevance", "correctness") and the canonical
    scorer name. This tool deliberately does not return a precomputed
    alias map — the agent uses the name + `doc` string to reason about fit
    and prompts the user for clarification when the match is ambiguous.

    Args:
        args: (ignored — kept for dispatch symmetry with other tools)

    Returns:
        dict: { scorers: [{name, doc}, ...] }
          - name: canonical MLflow scorer class name, pass this verbatim
            to submit_eval_job's `scorers` array.
          - doc:  first line of the class docstring (≤200 chars), use it
            to match user terminology to the canonical scorer.
    """
    from mlflow.genai import scorers as S  # type: ignore  # noqa: PLC0415

    # Any attribute on mlflow.genai.scorers that is a class and not private
    # is a scorer; that's the same convention the eval container uses to
    # populate SCORER_CATALOG. Filter accordingly.
    scorer_entries: list[dict] = []
    for name in sorted(dir(S)):
        if name.startswith("_"):
            continue
        obj = getattr(S, name)
        if not isinstance(obj, type):
            continue
        # The class lives in mlflow.genai.scorers; everything else is a helper.
        if not getattr(obj, "__module__", "").startswith("mlflow.genai.scorers"):
            continue
        scorer_entries.append({
            "name": name,
            "doc":  (obj.__doc__ or "").strip().split("\n", 1)[0][:200],
        })

    # QA BUG-002: trace-based retrieval scorers read retrieved chunks from
    # RETRIEVER spans on each row's trace. Static datasets evaluated via
    # prepare_eval_dataset have no live retriever, and in the SageMaker
    # Processing environment their judges return no metric. Annotate so the
    # agent steers "faithfulness"/groundedness requests to Guidelines (the
    # eval container instantiates it as `answer_groundedness`, judging the
    # answer against the context column passed via inputs).
    _TRACE_BASED = {"RetrievalGroundedness", "RetrievalRelevance", "RetrievalSufficiency"}
    for entry in scorer_entries:
        if entry["name"] in _TRACE_BASED:
            entry["doc"] += (
                " [NOTE: requires live RETRIEVER trace spans — for static "
                "datasets (prepare_eval_dataset) use Guidelines instead; it "
                "runs as answer_groundedness against the provided context.]"
            )
    return {"scorers": scorer_entries}


def _query_metrics(args: dict) -> dict:
    """Fetch metrics for an MLflow run.

    Args:
        args: Required: run_id.

    Returns:
        dict: metrics dict {metric_name: last_value}.
    """
    client = _mlflow_client()
    run = client.get_run(args["run_id"])
    return {"run_id": args["run_id"], "metrics": dict(run.data.metrics)}


def _retrieve_traces(args: dict) -> dict:
    """List recent runs for an experiment.

    Args:
        args: Required: experiment_name. Optional: max_results (default 10).

    Returns:
        dict: runs list with run_id, status, metrics summary.
    """
    client = _mlflow_client()
    exp = client.get_experiment_by_name(args["experiment_name"])
    if exp is None:
        raise RuntimeError(f"Experiment '{args['experiment_name']}' not found")
    runs = client.search_runs(
        experiment_ids=[exp.experiment_id],
        max_results=int(args.get("max_results", 10)),
    )
    return {
        "runs": [
            {"run_id": r.info.run_id, "status": r.info.status, "metrics": dict(r.data.metrics)}
            for r in runs
        ]
    }


def _analyze_trace(args: dict) -> dict:
    """Return full run data (params, metrics, tags) for a single MLflow run.

    Args:
        args: Required: run_id.

    Returns:
        dict: params, metrics, tags for the run.
    """
    client = _mlflow_client()
    run = client.get_run(args["run_id"])
    return {
        "run_id": args["run_id"],
        "params": dict(run.data.params),
        "metrics": dict(run.data.metrics),
        "tags": dict(run.data.tags),
        "status": run.info.status,
    }


# ── generate_compliance_report ─────────────────────────────────────────────

def _auto_failing_rows(run_id: str, n: int = _APPENDIX_ROWS) -> list[dict]:
    """Pull eval_results.parquet from the run and return the N worst rows.

    "Worst" = lowest aggregate mean across numeric scorer columns. If the
    artifact is missing (eval still running, or Processing job skipped it),
    returns an empty list rather than raising — the report is still useful.
    """
    import pandas as pd  # noqa: PLC0415

    client = _mlflow_client()
    try:
        with tempfile.TemporaryDirectory() as tmpdir:
            local = client.download_artifacts(
                run_id=run_id,
                path="eval_results.parquet",
                dst_path=tmpdir,
            )
            df = pd.read_parquet(local)
    except Exception as e:
        print(f"[compliance] no eval_results.parquet on run {run_id}: {e}")
        return []

    # Score columns land with names like "<scorer>/score" or "<scorer>/value".
    # Fall back to any float column if nothing matches that convention.
    score_cols = [c for c in df.columns if c.endswith("/score") or c.endswith("/value")]
    if not score_cols:
        score_cols = [c for c in df.columns if df[c].dtype.kind in "fi"]
    if not score_cols:
        return []

    df = df.copy()
    df["_agg"] = df[score_cols].mean(axis=1, numeric_only=True)
    worst = df.nsmallest(n, "_agg").drop(columns=["_agg"])
    # DataFrame.to_dict can emit numpy types that json.dumps trips on.
    return json.loads(worst.to_json(orient="records"))


def _render_markdown(*, report: dict, failing_rows: list[dict]) -> str:
    """Render the report dict + failing-rows appendix as markdown."""
    lines: list[str] = []
    lines.append(f"# {report.get('title', 'LLM Evaluation Compliance Report')}\n")
    lines.append(f"**Run ID:** `{report['run_id']}`  ")
    lines.append(f"**Generated:** {time.strftime('%Y-%m-%d %H:%M:%S UTC', time.gmtime())}\n")

    for section in report.get("sections", []):
        lines.append(f"## {section['heading']}\n")
        lines.append(f"{section['body']}\n")

    metrics = report.get("metrics") or {}
    if metrics:
        lines.append("## Metrics Summary\n")
        lines.append("| Metric | Value |")
        lines.append("|---|---|")
        for name, value in sorted(metrics.items()):
            lines.append(f"| `{name}` | {value} |")
        lines.append("")

    findings = report.get("findings") or []
    if findings:
        lines.append("## Findings\n")
        for f in findings:
            status = f.get("status", "?").upper()
            lines.append(f"- **[{status}] {f.get('name', '')}** — {f.get('body', '')}")
        lines.append("")

    if failing_rows:
        lines.append(f"## Appendix: Worst {len(failing_rows)} Failing Rows\n")
        lines.append("```json")
        lines.append(json.dumps(failing_rows, indent=2, default=str))
        lines.append("```")
    return "\n".join(lines)


def _render_docx(*, report: dict, failing_rows: list[dict]) -> bytes:
    """Render the report as a .docx with pass/fail coloring on Findings."""
    from docx import Document  # noqa: PLC0415
    from docx.shared import RGBColor  # noqa: PLC0415

    doc = Document()
    doc.add_heading(report.get("title", "LLM Evaluation Compliance Report"), level=0)
    meta = doc.add_paragraph()
    meta.add_run(f"Run ID: {report['run_id']}\n").bold = True
    meta.add_run(f"Generated: {time.strftime('%Y-%m-%d %H:%M:%S UTC', time.gmtime())}")

    for section in report.get("sections", []):
        doc.add_heading(section["heading"], level=1)
        doc.add_paragraph(section["body"])

    metrics = report.get("metrics") or {}
    if metrics:
        doc.add_heading("Metrics Summary", level=1)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Light List Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text = "Metric", "Value"
        for name, value in sorted(metrics.items()):
            row = table.add_row().cells
            row[0].text = str(name)
            row[1].text = str(value)

    findings = report.get("findings") or []
    if findings:
        doc.add_heading("Findings", level=1)
        for f in findings:
            status = f.get("status", "?").upper()
            para = doc.add_paragraph()
            run = para.add_run(f"[{status}] ")
            run.bold = True
            run.font.color.rgb = (
                RGBColor(0x1E, 0x87, 0x1E) if status == "PASS"
                else RGBColor(0xC4, 0x1E, 0x1E) if status == "FAIL"
                else RGBColor(0x80, 0x80, 0x80)
            )
            para.add_run(f"{f.get('name', '')} — {f.get('body', '')}")

    if failing_rows:
        doc.add_heading(f"Appendix: Worst {len(failing_rows)} Failing Rows", level=1)
        # One paragraph per row, JSON-stringified. Tables get unreadable fast
        # with variable eval row schemas, so this stays readable.
        for i, row in enumerate(failing_rows, start=1):
            doc.add_heading(f"Row {i}", level=2)
            doc.add_paragraph(json.dumps(row, indent=2, default=str))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _generate_compliance_report(args: dict) -> dict:
    """Render a compliance report and stamp the thread row with its S3 URI.

    Inputs from the caller are the *narrative* (title, sections, findings)
    authored by the agent under the `compliance-documentation` skill. This
    function is deterministic: it renders the markdown + .docx, uploads
    both to S3, grabs an auto-appendix of failing rows from the MLflow run,
    and writes `compliance_report_s3_uri` / `compliance_report_format` /
    `compliance_report_at` onto the thread row's eval job entry.

    Args:
        args: Required:
                thread_id, job_id, run_id, sections (list[{heading, body}]).
              Optional:
                title, findings (list[{name, status, body}]),
                metrics (dict), appendix_failing_rows (list[dict] — overrides
                the auto pull), formats (list[str] subset of ["md","docx"],
                default both), out_s3_prefix.

    Returns:
        dict: artifact_s3_uri (primary format), artifacts (per-format dict),
              appendix_rows (count), run_id, job_id, thread_id.
    """
    thread_id = args["thread_id"]
    job_id    = args["job_id"]
    run_id    = args["run_id"]
    sections  = args.get("sections", [])
    title     = args.get("title", "LLM Evaluation Compliance Report")
    findings  = args.get("findings") or []
    metrics   = args.get("metrics") or {}
    formats   = args.get("formats") or ["md", "docx"]
    out_prefix = (
        args.get("out_s3_prefix")
        or f"compliance-reports/{thread_id}/{job_id}"
    ).rstrip("/")

    appendix = args.get("appendix_failing_rows")
    if appendix is None:
        appendix = _auto_failing_rows(run_id)

    report = {
        "title":    title,
        "run_id":   run_id,
        "sections": sections,
        "findings": findings,
        "metrics":  metrics,
    }

    s3 = boto3.client("s3", region_name=AWS_REGION)
    artifacts: dict[str, str] = {}

    if "md" in formats:
        md_bytes = _render_markdown(report=report, failing_rows=appendix).encode("utf-8")
        md_key = f"{out_prefix}/report.md"
        s3.put_object(
            Bucket=SESSION_BUCKET, Key=md_key, Body=md_bytes,
            ContentType="text/markdown",
        )
        artifacts["md"] = f"s3://{SESSION_BUCKET}/{md_key}"

    if "docx" in formats:
        docx_bytes = _render_docx(report=report, failing_rows=appendix)
        docx_key = f"{out_prefix}/report.docx"
        s3.put_object(
            Bucket=SESSION_BUCKET, Key=docx_key, Body=docx_bytes,
            ContentType=(
                "application/vnd.openxmlformats-officedocument"
                ".wordprocessingml.document"
            ),
        )
        artifacts["docx"] = f"s3://{SESSION_BUCKET}/{docx_key}"

    if not artifacts:
        raise ValueError(f"formats must include 'md' or 'docx' (got {formats!r})")

    # Primary artifact the frontend download button points at — prefer docx.
    primary = artifacts.get("docx") or artifacts.get("md")
    primary_fmt = "docx" if "docx" in artifacts else "md"
    now = int(time.time())

    ddb = boto3.resource("dynamodb", region_name=AWS_REGION).Table(JOBS_TABLE)
    ddb.update_item(
        Key={"task_id": thread_id},
        UpdateExpression=(
            "SET jobs.#jid.compliance_report_s3_uri = :uri, "
            "    jobs.#jid.compliance_report_format = :fmt, "
            "    jobs.#jid.compliance_report_at     = :t, "
            "    jobs.#jid.updated_at               = :t, "
            "    updated_at                         = :t"
        ),
        ExpressionAttributeNames={"#jid": job_id},
        ExpressionAttributeValues={":uri": primary, ":fmt": primary_fmt, ":t": now},
    )

    return {
        "thread_id":        thread_id,
        "job_id":           job_id,
        "run_id":           run_id,
        "artifact_s3_uri":  primary,
        "artifact_format":  primary_fmt,
        "artifacts":        artifacts,
        "appendix_rows":    len(appendix),
    }


_DISPATCH: dict[str, Any] = {
    "list_scorers":               _list_scorers,
    "query_metrics":              _query_metrics,
    "retrieve_traces":            _retrieve_traces,
    "analyze_trace":              _analyze_trace,
    "generate_compliance_report": _generate_compliance_report,
}


def handler(event: dict, context: Any) -> dict:
    """Gateway MCP tool dispatcher for MLflow skills.

    Args:
        event: AgentCore Gateway passes the tool arguments map as `event`.
        context: Lambda context; tool name is in
            `context.client_context.custom['bedrockAgentCoreToolName']`,
            formatted as `${target_name}___${tool_name}`.

    Returns:
        dict: MCP content response.
    """
    raw_tool = context.client_context.custom.get("bedrockAgentCoreToolName", "") if getattr(context, "client_context", None) else ""
    tool_name = raw_tool.split("___", 1)[1] if "___" in raw_tool else raw_tool
    arguments = event or {}
    fn = _DISPATCH.get(tool_name)
    if fn is None:
        return {"content": [{"type": "text", "text": f"Unknown tool: {tool_name}"}], "isError": True}
    try:
        return {"content": [{"type": "text", "text": json.dumps(fn(arguments))}]}
    except Exception as e:
        return {"content": [{"type": "text", "text": f"Error: {e}"}], "isError": True}
